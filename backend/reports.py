import io
from datetime import datetime
from typing import Dict, Any, List
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

from backend.database import get_db_connection
from backend.settlement import calculate_trip_settlement

# Numbered canvas for "Page X of Y" in PDF
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(A4[0] - 36, 20, page_text)
        self.drawString(36, 20, "Spillter - Travel Group Expense Tracker & Settlement Statement")
        self.restoreState()


def get_trip_full_data(trip_id: int) -> Dict[str, Any]:
    settlement_data = calculate_trip_settlement(trip_id)
    conn = get_db_connection()
    cursor = conn.cursor()

    # Fetch expenses with payer and splits
    cursor.execute("""
        SELECT e.*, m.name as payer_name
        FROM expenses e
        JOIN members m ON e.payer_id = m.id
        WHERE e.trip_id = ?
        ORDER BY e.date ASC, e.id ASC
    """, (trip_id,))
    expense_rows = cursor.fetchall()

    expenses = []
    for exp_row in expense_rows:
        exp = dict(exp_row)
        # Fetch splits for this expense
        cursor.execute("""
            SELECT es.*, m.name as member_name
            FROM expense_splits es
            JOIN members m ON es.member_id = m.id
            WHERE es.expense_id = ?
        """, (exp["id"],))
        splits = [dict(s) for s in cursor.fetchall()]
        exp["splits"] = splits
        expenses.append(exp)

    conn.close()
    settlement_data["expenses"] = expenses
    return settlement_data


def generate_trip_pdf(trip_id: int) -> io.BytesIO:
    data = get_trip_full_data(trip_id)
    trip = data["trip"]
    currency = trip.get("currency", "₹")
    # For PDF, use standard symbols or text representation to avoid encoding issues
    curr_symbol = "Rs. " if currency == "₹" else f"{currency} "

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1E1B4B"),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#475569"),
        spaceAfter=12
    )
    h2_style = ParagraphStyle(
        'Heading2Custom',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#312E81"),
        spaceBefore=14,
        spaceAfter=6
    )
    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1F2937")
    )
    cell_bold_style = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1F2937")
    )
    header_cell_style = ParagraphStyle(
        'HeaderCell',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.white
    )

    amount_cell_style = ParagraphStyle(
        'AmountCell',
        parent=cell_bold_style,
        alignment=2
    )
    header_right_style = ParagraphStyle(
        'HeaderRight',
        parent=header_cell_style,
        alignment=2
    )

    story = []

    # Title & Header
    # Remove non-ascii emojis for standard PDF font compatibility
    clean_trip_name = trip['name'].encode('ascii', 'ignore').decode().strip() or trip['name']
    clean_desc = (trip.get('description') or '').encode('ascii', 'ignore').decode().strip()

    story.append(Paragraph(f"Trip Expense Statement: {clean_trip_name}", title_style))
    if clean_desc:
        story.append(Paragraph(clean_desc, subtitle_style))
    
    gen_time = datetime.now().strftime("%d %b %Y, %I:%M %p")
    story.append(Paragraph(f"<b>Generated on:</b> {gen_time} | <b>Currency:</b> {trip.get('currency', '₹')}", cell_style))
    story.append(Spacer(1, 10))

    # Summary KPI Table
    kpi_data = [
        [
            Paragraph("<b>Total Trip Expense</b>", cell_style),
            Paragraph("<b>Total Members</b>", cell_style),
            Paragraph("<b>Average / Person</b>", cell_style),
            Paragraph("<b>Total Logged Bills</b>", cell_style)
        ],
        [
            Paragraph(f"<b>{curr_symbol}{data['total_expenses']:,.2f}</b>", ParagraphStyle('Kpi1', parent=cell_style, fontSize=12, textColor=colors.HexColor("#1E1B4B"))),
            Paragraph(f"<b>{data['member_count']} Persons</b>", ParagraphStyle('Kpi2', parent=cell_style, fontSize=12, textColor=colors.HexColor("#1E1B4B"))),
            Paragraph(f"<b>{curr_symbol}{data['average_expense']:,.2f}</b>", ParagraphStyle('Kpi3', parent=cell_style, fontSize=12, textColor=colors.HexColor("#1E1B4B"))),
            Paragraph(f"<b>{len(data['expenses'])} Items</b>", ParagraphStyle('Kpi4', parent=cell_style, fontSize=12, textColor=colors.HexColor("#1E1B4B")))
        ]
    ]
    t_kpi = Table(kpi_data, colWidths=[130, 130, 130, 130])
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor("#EEF2FF")),
        ('BACKGROUND', (1, 0), (1, -1), colors.HexColor("#F8FAFC")),
        ('BACKGROUND', (2, 0), (2, -1), colors.HexColor("#ECFDF5")),
        ('BACKGROUND', (3, 0), (3, -1), colors.HexColor("#FEF3C7")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t_kpi)
    story.append(Spacer(1, 14))

    # Section: Simplified Settlements ("Who Pays Whom")
    story.append(Paragraph("1. Final Settlement Summary (Who Pays Whom)", h2_style))
    story.append(Paragraph("Debt-minimization algorithm result: settle all group balances in the minimum transactions.", cell_style))
    story.append(Spacer(1, 4))

    if data["settlements"]:
        settle_table_data = [
            [
                Paragraph("<b>From (Payer)</b>", header_cell_style),
                Paragraph("<b>Action</b>", header_cell_style),
                Paragraph("<b>To (Receiver)</b>", header_cell_style),
                Paragraph("<b>Amount to Pay</b>", header_right_style)
            ]
        ]
        for s in data["settlements"]:
            settle_table_data.append([
                Paragraph(f"<b>{s['from_name']}</b>", cell_style),
                Paragraph("pays ➔", cell_style),
                Paragraph(f"<b>{s['to_name']}</b>", cell_style),
                Paragraph(f"<b>{curr_symbol}{s['amount']:,.2f}</b>", amount_cell_style)
            ])
        t_settle = Table(settle_table_data, colWidths=[150, 80, 150, 140])
        t_settle.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ('PADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(t_settle)
    else:
        story.append(Paragraph("All members are completely settled up.", cell_style))

    if data.get("settlement_history"):
        story.append(Spacer(1, 10))
        story.append(Paragraph("<b>Completed Settlement Payments:</b>", cell_bold_style))
        story.append(Spacer(1, 4))
        comp_data = [
            [
                Paragraph("<b>Date</b>", header_cell_style),
                Paragraph("<b>Payer</b>", header_cell_style),
                Paragraph("<b>Receiver</b>", header_cell_style),
                Paragraph("<b>Amount Paid</b>", header_right_style)
            ]
        ]
        for sh in data["settlement_history"]:
            raw_sh_date = sh.get("date", "")
            try:
                disp_sh_date = datetime.strptime(raw_sh_date, "%Y-%m-%d").strftime("%d %b %Y")
            except Exception:
                disp_sh_date = raw_sh_date
            comp_data.append([
                Paragraph(disp_sh_date, cell_style),
                Paragraph(sh["payer_name"], cell_style),
                Paragraph(sh["receiver_name"], cell_style),
                Paragraph(f"<b>{curr_symbol}{float(sh['amount']):,.2f}</b>", amount_cell_style)
            ])
        t_comp = Table(comp_data, colWidths=[90, 150, 150, 130])
        t_comp.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#059669")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_comp)

    story.append(Spacer(1, 14))

    # Section: Member Balance Breakdown
    story.append(Paragraph("2. Member Balance Summary", h2_style))
    mem_table_data = [
        [
            Paragraph("<b>Member</b>", header_cell_style),
            Paragraph("<b>Total Paid</b>", header_cell_style),
            Paragraph("<b>Total Share (Owed)</b>", header_cell_style),
            Paragraph("<b>Net Balance</b>", header_cell_style),
            Paragraph("<b>Status</b>", header_cell_style),
        ]
    ]
    for m in data["member_stats"]:
        net = m["net_balance"]
        if net > 0.009:
            status = "Gets Back"
            net_text = f"+{curr_symbol}{abs(net):,.2f}"
            color = colors.HexColor("#065F46") # Emerald dark
        elif net < -0.009:
            status = "Owes"
            net_text = f"-{curr_symbol}{abs(net):,.2f}"
            color = colors.HexColor("#991B1B") # Red dark
        else:
            status = "Settled"
            net_text = f"{curr_symbol}0.00"
            color = colors.HexColor("#475569")

        mem_table_data.append([
            Paragraph(f"<b>{m['name']}</b>", cell_style),
            Paragraph(f"{curr_symbol}{m['total_paid']:,.2f}", cell_style),
            Paragraph(f"{curr_symbol}{m['total_owed']:,.2f}", cell_style),
            Paragraph(f"<b>{net_text}</b>", ParagraphStyle('NetStyle', parent=cell_style, textColor=color)),
            Paragraph(status, ParagraphStyle('StatStyle', parent=cell_style, textColor=color, fontName='Helvetica-Bold')),
        ])

    t_mem = Table(mem_table_data, colWidths=[120, 100, 100, 100, 100])
    t_mem.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#334155")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('PADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t_mem)
    story.append(Spacer(1, 14))

    # Section: Itemized Expenses Log Statement
    story.append(Paragraph("3. Itemized Expense Statement (All Entries)", h2_style))
    exp_table_data = [
        [
            Paragraph("<b>Date</b>", header_cell_style),
            Paragraph("<b>Expense Title</b>", header_cell_style),
            Paragraph("<b>Category</b>", header_cell_style),
            Paragraph("<b>Paid By</b>", header_cell_style),
            Paragraph("<b>Mode</b>", header_cell_style),
            Paragraph("<b>Split Status</b>", header_cell_style),
            Paragraph("<b>Amount</b>", header_right_style),
        ]
    ]

    total_member_count = data.get("member_count", 1)
    for e in data["expenses"]:
        clean_title = e["title"].encode('ascii', 'ignore').decode().strip() or e["title"]
        clean_notes = (e.get("notes") or "").encode('ascii', 'ignore').decode().strip()
        title_html = f"<b>{clean_title}</b>"
        if clean_notes:
            title_html += f"<br/><font color='#64748B' size='7'><i>{clean_notes}</i></font>"

        mode = e.get("payment_mode") or "UPI"
        raw_date = e.get("date", "")
        try:
            display_date = datetime.strptime(raw_date, "%Y-%m-%d").strftime("%d %b %Y")
        except Exception:
            display_date = raw_date

        splits = e.get("splits", [])
        split_count = len(splits)
        split_type = (e.get("split_type") or "equal").lower()

        if split_count >= total_member_count and split_type == "equal":
            split_html = "<b>Equal</b><br/><font color='#059669' size='7'>All members</font>"
        elif split_count < total_member_count:
            names = ", ".join(s.get("member_name", "") for s in splits)
            split_html = f"<b>{split_count} Members</b><br/><font color='#4F46E5' size='7'>{names}</font>"
        else:
            shares = ", ".join(f"{s.get('member_name', '')[:5]}:{curr_symbol}{s.get('share_amount', 0):,.0f}" for s in splits)
            split_html = f"<b>Custom</b><br/><font color='#D97706' size='7'>{shares}</font>"

        exp_table_data.append([
            Paragraph(f"<b>{display_date}</b>", cell_style),
            Paragraph(title_html, cell_style),
            Paragraph(e["category"], cell_style),
            Paragraph(f"<b>{e['payer_name']}</b>", cell_style),
            Paragraph(mode, cell_style),
            Paragraph(split_html, cell_style),
            Paragraph(f"<b>{curr_symbol}{e['amount']:,.2f}</b>", amount_cell_style),
        ])

    t_exp = Table(exp_table_data, colWidths=[65, 115, 55, 65, 45, 105, 73])
    t_exp.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 4.5),
    ]))
    story.append(t_exp)

    # Build PDF
    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


def get_member_full_data(trip_id: int, member_id: int) -> Dict[str, Any]:
    settlement_data = calculate_trip_settlement(trip_id)
    trip = settlement_data["trip"]

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM members WHERE id = ? AND trip_id = ?", (member_id, trip_id))
    member_row = cursor.fetchone()
    if not member_row:
        conn.close()
        raise ValueError("Member not found in trip")
    member = dict(member_row)

    # 1. Expenses PAID BY this member
    cursor.execute("""
        SELECT e.*, m.name as payer_name
        FROM expenses e
        JOIN members m ON e.payer_id = m.id
        WHERE e.trip_id = ? AND e.payer_id = ?
        ORDER BY e.date ASC, e.id ASC
    """, (trip_id, member_id))
    paid_expenses = []
    for r in cursor.fetchall():
        exp = dict(r)
        cursor.execute("""
            SELECT es.*, m.name as member_name
            FROM expense_splits es
            JOIN members m ON es.member_id = m.id
            WHERE es.expense_id = ?
        """, (exp["id"],))
        exp["splits"] = [dict(s) for s in cursor.fetchall()]
        paid_expenses.append(exp)

    # 2. Expenses SHARED BY this member (where someone else was the upfront payer)
    cursor.execute("""
        SELECT e.*, m.name as payer_name, es.share_amount as member_share, es.percentage as member_percentage
        FROM expense_splits es
        JOIN expenses e ON es.expense_id = e.id
        JOIN members m ON e.payer_id = m.id
        WHERE e.trip_id = ? AND es.member_id = ? AND e.payer_id != ?
        ORDER BY e.date ASC, e.id ASC
    """, (trip_id, member_id, member_id))
    shared_expenses = [dict(r) for r in cursor.fetchall()]

    # 3. Direct settlements history involving this member
    cursor.execute("""
        SELECT sp.*, p.name as payer_name, r.name as receiver_name
        FROM settlement_payments sp
        JOIN members p ON sp.payer_id = p.id
        JOIN members r ON sp.receiver_id = r.id
        WHERE sp.trip_id = ? AND (sp.payer_id = ? OR sp.receiver_id = ?)
        ORDER BY sp.date DESC, sp.id DESC
    """, (trip_id, member_id, member_id))
    settlement_history = [dict(r) for r in cursor.fetchall()]

    conn.close()

    member_stat = next((m for m in settlement_data["member_stats"] if m["id"] == member_id), None)
    member_settlements_to_pay = [s for s in settlement_data["settlements"] if s["from_id"] == member_id]
    member_settlements_to_receive = [s for s in settlement_data["settlements"] if s["to_id"] == member_id]

    return {
        "trip": trip,
        "member": member,
        "member_stat": member_stat,
        "paid_expenses": paid_expenses,
        "shared_expenses": shared_expenses,
        "settlement_history": settlement_history,
        "settlements_to_pay": member_settlements_to_pay,
        "settlements_to_receive": member_settlements_to_receive,
        "all_members_count": settlement_data["member_count"]
    }


def generate_member_pdf(trip_id: int, member_id: int) -> io.BytesIO:
    data = get_member_full_data(trip_id, member_id)
    trip = data["trip"]
    member = data["member"]
    currency = trip.get("currency", "₹")
    curr_symbol = "Rs. " if currency == "₹" else f"{currency} "

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1E1B4B"),
        spaceAfter=3
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#475569"),
        spaceAfter=10
    )
    h2_style = ParagraphStyle(
        'Heading2Custom',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#312E81"),
        spaceBefore=12,
        spaceAfter=5
    )
    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#1F2937")
    )
    cell_bold_style = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#1F2937")
    )
    amount_cell_style = ParagraphStyle(
        'AmountCell',
        parent=cell_bold_style,
        alignment=2
    )
    header_cell_style = ParagraphStyle(
        'HeaderCell',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white
    )
    header_right_style = ParagraphStyle(
        'HeaderRight',
        parent=header_cell_style,
        alignment=2
    )

    story = []

    clean_trip_name = trip['name'].encode('ascii', 'ignore').decode().strip() or trip['name']
    clean_member_name = member['name'].encode('ascii', 'ignore').decode().strip() or member['name']
    member_phone = member.get('phone') or 'Not provided'

    # Title Banner
    story.append(Paragraph(f"Individual Traveler Statement: {clean_member_name}", title_style))
    gen_time = datetime.now().strftime("%d %b %Y, %I:%M %p")
    story.append(Paragraph(f"Trip: <b>{clean_trip_name}</b> | Phone: {member_phone} | Generated: {gen_time} | Currency: {trip.get('currency', '₹')}", subtitle_style))
    story.append(Spacer(1, 4))

    # KPI Summary Card for this member
    m_stat = data.get("member_stat") or {"total_paid": 0.0, "total_owed": 0.0, "net_balance": 0.0}
    net_val = m_stat.get("net_balance", 0.0)
    if net_val > 0.009:
        net_status = "Gets Back (To Receive)"
        net_str = f"+{curr_symbol}{abs(net_val):,.2f}"
        net_color = colors.HexColor("#065F46")
        net_bg = colors.HexColor("#ECFDF5")
    elif net_val < -0.009:
        net_status = "Owes (To Pay)"
        net_str = f"-{curr_symbol}{abs(net_val):,.2f}"
        net_color = colors.HexColor("#991B1B")
        net_bg = colors.HexColor("#FEF2F2")
    else:
        net_status = "All Settled Up"
        net_str = f"{curr_symbol}0.00"
        net_color = colors.HexColor("#334155")
        net_bg = colors.HexColor("#F1F5F9")

    kpi_data = [
        [
            Paragraph("<b>Total Paid Upfront</b>", cell_style),
            Paragraph("<b>Fair Share (Consumed)</b>", cell_style),
            Paragraph(f"<b>Net Position ({net_status})</b>", cell_style),
            Paragraph("<b>Associated Bills</b>", cell_style)
        ],
        [
            Paragraph(f"<b>{curr_symbol}{m_stat['total_paid']:,.2f}</b>", ParagraphStyle('KP1', parent=cell_style, fontSize=11, fontName='Helvetica-Bold', textColor=colors.HexColor("#1E1B4B"))),
            Paragraph(f"<b>{curr_symbol}{m_stat['total_owed']:,.2f}</b>", ParagraphStyle('KP2', parent=cell_style, fontSize=11, fontName='Helvetica-Bold', textColor=colors.HexColor("#1E1B4B"))),
            Paragraph(f"<b>{net_str}</b>", ParagraphStyle('KP3', parent=cell_style, fontSize=11, fontName='Helvetica-Bold', textColor=net_color)),
            Paragraph(f"<b>{len(data['paid_expenses']) + len(data['shared_expenses'])} Items</b>", ParagraphStyle('KP4', parent=cell_style, fontSize=11, fontName='Helvetica-Bold', textColor=colors.HexColor("#1E1B4B")))
        ]
    ]
    t_kpi = Table(kpi_data, colWidths=[130, 130, 150, 113])
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor("#EEF2FF")),
        ('BACKGROUND', (1, 0), (1, -1), colors.HexColor("#F8FAFC")),
        ('BACKGROUND', (2, 0), (2, -1), net_bg),
        ('BACKGROUND', (3, 0), (3, -1), colors.HexColor("#F1F5F9")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('PADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t_kpi)
    story.append(Spacer(1, 10))

    # Section 1: Settlements for this Member
    story.append(Paragraph(f"1. Direct Settlement Instructions for {clean_member_name}", h2_style))
    settle_rows = []
    if data["settlements_to_receive"]:
        for s in data["settlements_to_receive"]:
            settle_rows.append([
                Paragraph(f"<b>{s['from_name']}</b>", cell_style),
                Paragraph("owes and pays ➔", cell_style),
                Paragraph(f"<b>{clean_member_name}</b>", cell_style),
                Paragraph(f"<b>+{curr_symbol}{s['amount']:,.2f}</b>", ParagraphStyle('GreenA', parent=amount_cell_style, textColor=colors.HexColor("#065F46")))
            ])
    if data["settlements_to_pay"]:
        for s in data["settlements_to_pay"]:
            settle_rows.append([
                Paragraph(f"<b>{clean_member_name}</b>", cell_style),
                Paragraph("owes and pays ➔", cell_style),
                Paragraph(f"<b>{s['to_name']}</b>", cell_style),
                Paragraph(f"<b>-{curr_symbol}{s['amount']:,.2f}</b>", ParagraphStyle('RedA', parent=amount_cell_style, textColor=colors.HexColor("#991B1B")))
            ])

    if settle_rows:
        s_table_data = [[
            Paragraph("<b>Payer</b>", header_cell_style),
            Paragraph("<b>Action</b>", header_cell_style),
            Paragraph("<b>Receiver</b>", header_cell_style),
            Paragraph("<b>Amount</b>", header_right_style),
        ]] + settle_rows
        t_s = Table(s_table_data, colWidths=[150, 90, 150, 133])
        t_s.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_s)
    else:
        story.append(Paragraph("No pending settlements. This traveler is completely settled up!", cell_style))

    story.append(Spacer(1, 10))

    # Section 2: Expenses PAID UPFRONT by this member
    story.append(Paragraph(f"2. Expenses Paid Upfront by {clean_member_name}", h2_style))
    if data["paid_expenses"]:
        paid_table_data = [[
            Paragraph("<b>Date</b>", header_cell_style),
            Paragraph("<b>Expense Title</b>", header_cell_style),
            Paragraph("<b>Category</b>", header_cell_style),
            Paragraph("<b>Mode</b>", header_cell_style),
            Paragraph("<b>Split With</b>", header_cell_style),
            Paragraph("<b>Total Bill</b>", header_right_style),
            Paragraph("<b>Own Share</b>", header_right_style),
            Paragraph("<b>To Recover</b>", header_right_style),
        ]]
        total_paid_sum = 0.0
        total_recover_sum = 0.0
        for pe in data["paid_expenses"]:
            raw_date = pe.get("date", "")
            try:
                display_date = datetime.strptime(raw_date, "%Y-%m-%d").strftime("%d %b %Y")
            except Exception:
                display_date = raw_date

            clean_title = pe["title"].encode('ascii', 'ignore').decode().strip() or pe["title"]
            mode = pe.get("payment_mode") or "UPI"
            splits = pe.get("splits", [])
            own_s = next((s for s in splits if s["member_id"] == member_id), None)
            own_share = own_s["share_amount"] if own_s else 0.0
            to_recover = pe["amount"] - own_share
            total_paid_sum += pe["amount"]
            total_recover_sum += to_recover

            names = ", ".join(s["member_name"] for s in splits)
            paid_table_data.append([
                Paragraph(f"<b>{display_date}</b>", cell_style),
                Paragraph(clean_title, cell_style),
                Paragraph(pe["category"], cell_style),
                Paragraph(mode, cell_style),
                Paragraph(f"<font size='7' color='#475569'>{names}</font>", cell_style),
                Paragraph(f"{curr_symbol}{pe['amount']:,.2f}", amount_cell_style),
                Paragraph(f"{curr_symbol}{own_share:,.2f}", amount_cell_style),
                Paragraph(f"<b>+{curr_symbol}{to_recover:,.2f}</b>", ParagraphStyle('Rec', parent=amount_cell_style, textColor=colors.HexColor("#065F46"))),
            ])

        paid_table_data.append([
            Paragraph("<b>TOTALS</b>", cell_bold_style),
            Paragraph("", cell_style),
            Paragraph("", cell_style),
            Paragraph("", cell_style),
            Paragraph("", cell_style),
            Paragraph(f"<b>{curr_symbol}{total_paid_sum:,.2f}</b>", amount_cell_style),
            Paragraph("", cell_style),
            Paragraph(f"<b>+{curr_symbol}{total_recover_sum:,.2f}</b>", ParagraphStyle('TotRec', parent=amount_cell_style, textColor=colors.HexColor("#065F46"))),
        ])

        t_paid = Table(paid_table_data, colWidths=[58, 105, 50, 42, 100, 56, 56, 56])
        t_paid.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#F1F5F9")),
            ('PADDING', (0, 0), (-1, -1), 3.5),
        ]))
        story.append(t_paid)
    else:
        story.append(Paragraph(f"No expenses paid upfront by {clean_member_name}.", cell_style))

    story.append(Spacer(1, 10))

    # Section 3: Expenses SHARED by this member (paid by others)
    story.append(Paragraph(f"3. Expenses Shared by {clean_member_name} (Paid by Other Travelers)", h2_style))
    if data["shared_expenses"]:
        shared_table_data = [[
            Paragraph("<b>Date</b>", header_cell_style),
            Paragraph("<b>Expense Title</b>", header_cell_style),
            Paragraph("<b>Category</b>", header_cell_style),
            Paragraph("<b>Paid Upfront By</b>", header_cell_style),
            Paragraph("<b>Total Bill</b>", header_right_style),
            Paragraph("<b>Share Owed</b>", header_right_style),
        ]]
        total_owed_sum = 0.0
        for se in data["shared_expenses"]:
            raw_date = se.get("date", "")
            try:
                display_date = datetime.strptime(raw_date, "%Y-%m-%d").strftime("%d %b %Y")
            except Exception:
                display_date = raw_date

            clean_title = se["title"].encode('ascii', 'ignore').decode().strip() or se["title"]
            share_val = se.get("member_share", 0.0)
            total_owed_sum += share_val

            shared_table_data.append([
                Paragraph(f"<b>{display_date}</b>", cell_style),
                Paragraph(clean_title, cell_style),
                Paragraph(se["category"], cell_style),
                Paragraph(f"<b>{se['payer_name']}</b>", cell_style),
                Paragraph(f"{curr_symbol}{se['amount']:,.2f}", amount_cell_style),
                Paragraph(f"<b>-{curr_symbol}{share_val:,.2f}</b>", ParagraphStyle('OweVal', parent=amount_cell_style, textColor=colors.HexColor("#991B1B"))),
            ])

        shared_table_data.append([
            Paragraph("<b>TOTAL OWED</b>", cell_bold_style),
            Paragraph("", cell_style),
            Paragraph("", cell_style),
            Paragraph("", cell_style),
            Paragraph("", cell_style),
            Paragraph(f"<b>-{curr_symbol}{total_owed_sum:,.2f}</b>", ParagraphStyle('TotOwe', parent=amount_cell_style, textColor=colors.HexColor("#991B1B"))),
        ])

        t_shared = Table(shared_table_data, colWidths=[65, 138, 65, 95, 80, 80])
        t_shared.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#334155")),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#F1F5F9")),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_shared)
    else:
        story.append(Paragraph(f"No shared expenses where others paid for {clean_member_name}.", cell_style))

    # Build PDF
    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


from openpyxl import Workbook

def generate_trip_excel(trip_id: int) -> io.BytesIO:
    data = get_trip_full_data(trip_id)
    trip = data["trip"]
    currency = trip.get("currency", "₹")

    output = io.BytesIO()
    workbook = Workbook()

    # Define Styles
    header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E1B4B")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="475569")
    bold_font = Font(name="Calibri", size=11, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    currency_format = f'"{currency}" #,##0.00'

    # --- SHEET 1: Summary & Settlements ---
    ws_summary = workbook.active
    ws_summary.title = "Summary & Settlements"

    ws_summary["A1"] = f"Trip Statement: {trip['name']}"
    ws_summary["A1"].font = title_font
    ws_summary["A2"] = f"{trip.get('description', '')} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws_summary["A2"].font = subtitle_font

    # Overview Box
    ws_summary["A4"] = "Total Trip Expense"
    ws_summary["B4"] = data["total_expenses"]
    ws_summary["B4"].number_format = currency_format
    ws_summary["B4"].font = bold_font

    ws_summary["A5"] = "Total Members"
    ws_summary["B5"] = data["member_count"]

    ws_summary["A6"] = "Average Expense / Member"
    ws_summary["B6"] = data["average_expense"]
    ws_summary["B6"].number_format = currency_format
    ws_summary["B6"].font = bold_font

    ws_summary["A7"] = "Total Expense Entries"
    ws_summary["B7"] = len(data["expenses"])

    for r in range(4, 8):
        ws_summary[f"A{r}"].font = bold_font
        ws_summary[f"A{r}"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        ws_summary[f"A{r}"].border = thin_border
        ws_summary[f"B{r}"].border = thin_border

    # Member Balances Table
    ws_summary["A9"] = "MEMBER BALANCE SUMMARY"
    ws_summary["A9"].font = Font(name="Calibri", size=12, bold=True, color="312E81")

    mem_headers = ["Member Name", "Total Paid", "Total Share (Owed)", "Net Balance", "Status"]
    for col_idx, h in enumerate(mem_headers, start=1):
        cell = ws_summary.cell(row=10, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    row_idx = 11
    for m in data["member_stats"]:
        ws_summary.cell(row=row_idx, column=1, value=m["name"]).border = thin_border
        
        c_paid = ws_summary.cell(row=row_idx, column=2, value=m["total_paid"])
        c_paid.number_format = currency_format
        c_paid.border = thin_border

        c_owed = ws_summary.cell(row=row_idx, column=3, value=m["total_owed"])
        c_owed.number_format = currency_format
        c_owed.border = thin_border

        c_net = ws_summary.cell(row=row_idx, column=4, value=m["net_balance"])
        c_net.number_format = currency_format
        c_net.border = thin_border
        c_net.font = bold_font

        status_val = "Gets Back" if m["net_balance"] > 0.009 else ("Owes" if m["net_balance"] < -0.009 else "Settled")
        c_stat = ws_summary.cell(row=row_idx, column=5, value=status_val)
        c_stat.border = thin_border
        c_stat.alignment = Alignment(horizontal="center")
        if status_val == "Gets Back":
            c_stat.font = Font(name="Calibri", color="047857", bold=True)
        elif status_val == "Owes":
            c_stat.font = Font(name="Calibri", color="B91C1C", bold=True)
        row_idx += 1

    # Settlements Table ("Who Pays Whom")
    row_idx += 2
    ws_summary.cell(row=row_idx, column=1, value="FINAL SETTLEMENTS (WHO PAYS WHOM)").font = Font(name="Calibri", size=12, bold=True, color="312E81")
    row_idx += 1

    set_headers = ["Payer (Debtor)", "Action", "Receiver (Creditor)", "Amount to Pay"]
    for col_idx, h in enumerate(set_headers, start=1):
        cell = ws_summary.cell(row=row_idx, column=col_idx, value=h)
        cell.fill = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
    row_idx += 1

    if data["settlements"]:
        for s in data["settlements"]:
            ws_summary.cell(row=row_idx, column=1, value=s["from_name"]).border = thin_border
            act = ws_summary.cell(row=row_idx, column=2, value="pays ➔")
            act.alignment = Alignment(horizontal="center")
            act.border = thin_border
            ws_summary.cell(row=row_idx, column=3, value=s["to_name"]).border = thin_border
            
            amt_cell = ws_summary.cell(row=row_idx, column=4, value=s["amount"])
            amt_cell.number_format = currency_format
            amt_cell.font = bold_font
            amt_cell.border = thin_border
            row_idx += 1
    else:
        ws_summary.cell(row=row_idx, column=1, value="All expenses are completely settled!").border = thin_border

    # Adjust summary column widths
    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # --- SHEET 2: Detailed Expense Statement ---
    ws_expenses = workbook.create_sheet(title="All Expenses Statement")
    ws_expenses["A1"] = f"Detailed Expense Statement - {trip['name']}"
    ws_expenses["A1"].font = title_font

    exp_headers = ["ID", "Date", "Title", "Category", "Paid By", "Payment Mode", "Split Type", "Amount", "Split Breakdown", "Notes"]
    for col_idx, h in enumerate(exp_headers, start=1):
        cell = ws_expenses.cell(row=3, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    row_idx = 4
    for e in data["expenses"]:
        splits_summary = ", ".join([f"{s['member_name']} ({currency}{s['share_amount']:,.2f})" for s in e["splits"]])
        ws_expenses.cell(row=row_idx, column=1, value=e["id"]).border = thin_border
        ws_expenses.cell(row=row_idx, column=2, value=e["date"]).border = thin_border
        ws_expenses.cell(row=row_idx, column=3, value=e["title"]).border = thin_border
        ws_expenses.cell(row=row_idx, column=4, value=e["category"]).border = thin_border
        ws_expenses.cell(row=row_idx, column=5, value=e["payer_name"]).border = thin_border
        ws_expenses.cell(row=row_idx, column=6, value=e.get("payment_mode") or "UPI").border = thin_border
        ws_expenses.cell(row=row_idx, column=7, value=e["split_type"].capitalize()).border = thin_border

        amt_cell = ws_expenses.cell(row=row_idx, column=8, value=e["amount"])
        amt_cell.number_format = currency_format
        amt_cell.font = bold_font
        amt_cell.border = thin_border

        ws_expenses.cell(row=row_idx, column=9, value=splits_summary).border = thin_border
        ws_expenses.cell(row=row_idx, column=10, value=e.get("notes") or "").border = thin_border
        row_idx += 1

    # Total row
    ws_expenses.cell(row=row_idx, column=7, value="Total").font = bold_font
    ws_expenses.cell(row=row_idx, column=7).border = thin_border
    total_cell = ws_expenses.cell(row=row_idx, column=8, value=data["total_expenses"])
    total_cell.number_format = currency_format
    total_cell.font = bold_font
    total_cell.border = thin_border

    for col in ws_expenses.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_expenses.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    workbook.save(output)
    output.seek(0)
    return output
